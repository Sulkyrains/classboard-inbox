"""生成「知可而办 · 班级通知看板 发布说明.docx」，用于发到班级群里。

用法：python scripts/make-release-doc.py [输出目录]
依赖：python-docx（pip install python-docx）
注意：文档里不写任何明文密码，初始密码由班长另行分发。
"""
import os
import sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Pt, RGBColor

SITE = 'https://classboard-upc.pages.dev/'
APK = 'https://classboard-upc.pages.dev/classboard.apk'
VERSION = '2.3.2'
VERSION_CODE = 12
UPDATED = '2026-09-24'
FILENAME = '知可而办 · 班级通知看板 发布说明.docx'
FONT = '微软雅黑'


def set_font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_hyperlink(paragraph, url, text):
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), r_id)
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    props.append(OxmlElement('w:rFonts'))
    props[-1].set(qn('w:eastAsia'), FONT)
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '2563EB')
    props.append(color)
    underline = OxmlElement('w:u')
    underline.set(qn('w:val'), 'single')
    props.append(underline)
    run.append(props)
    text_node = OxmlElement('w:t')
    text_node.text = text
    run.append(text_node)
    link.append(run)
    paragraph._p.append(link)


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_font(run, size={1: 15, 2: 13, 3: 12}.get(level, 12), bold=True, color='1F3864')
    return p


def paragraph(doc, text, size=10.5, bold=False, color=None, style=None):
    p = doc.add_paragraph(style=style)
    set_font(p.add_run(text), size=size, bold=bold, color=color)
    return p


def bullet(doc, text, size=10.5):
    p = doc.add_paragraph(style='List Bullet')
    set_font(p.add_run(text), size=size)
    return p


def build(out_dir):
    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Pt(50)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(title.add_run('知可而办 · 班级通知看板'), size=20, bold=True, color='1F3864')
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(sub.add_run(f'发布说明 · {UPDATED} · v{VERSION}'), size=10.5, color='6B7280')

    paragraph(doc, '班级的作业、活动、班会和课程调整统一发在这里，一眼看清「什么时候、在哪、什么时候截止」，'
                   '再也不用在群里往上翻聊天记录。', size=11)
    doc.add_paragraph()

    heading(doc, '一、网址与下载', 1)
    p = doc.add_paragraph()
    set_font(p.add_run('网页版（电脑、手机都能开）：'), size=10.5)
    add_hyperlink(p, SITE, SITE)
    p = doc.add_paragraph()
    set_font(p.add_run('安卓 App（推荐，能收手机通知）：'), size=10.5)
    add_hyperlink(p, APK, APK)
    paragraph(doc, '也可以打开网页后点「下载安卓 App」按钮。iPhone / iPad 用 Safari 打开网页，'
                   '点底部「分享」→「添加到主屏幕」，就能像 App 一样使用。', size=10.5)

    heading(doc, '二、第一次使用（三步）', 1)
    for text in [
        '登录：用户名填学号，初始密码请私聊班长获取；第一次登录会要求改成只有自己知道的密码。',
        '打开通知开关：进入「个人账号 → 手机通知」，按提示允许通知（安卓 App 首次打开也会自动请求）。',
        '保持后台能收到：在手机设置里把本应用加入「自启动 / 允许后台活动」白名单，并把电池优化设为「不限制」。'
        'App 内「个人账号」页有自检面板，可看到后台检查次数在增长。',
    ]:
        bullet(doc, text)
    paragraph(doc, '小提示：下课后从后台把 App 划掉，通知就会停；遇到收不到通知，先看自启动和电池优化。', size=10)

    heading(doc, '三、这版更新了什么', 1)
    for text in [
        '日程更清楚：一条通知在日程里只出现一次，「活动时间」更名为「开始时间」，卡片里同时写明截止时间；'
        '已经点过「已完成」的通知不再占日程；左侧日期显示「今天 / 明天 / 后天」。',
        '日历订阅、手机通知里的时间也更规范：开始与截止分开标注，不会再看混。',
        '班委误操作可以撤销：已归档的通知点「撤销归档」会重新发布；已驳回的草稿点「撤销驳回」回到待审核。',
        '群消息导入更强：写得比较随意的消息也能生成草稿，没认出来的片段会列出来，可一键转成草稿人工补齐。',
        '时间提示更聪明：年份、日期明显不对时会提示核对，避免通知写错时间。',
    ]:
        bullet(doc, text)

    heading(doc, '四、主要功能', 1)
    for text in [
        '四种分类（作业 / 活动 / 事务 / 课程）与置顶，支持按分类、时间、未读筛选和关键词搜索。',
        '今日看板：今天要交的、要开的、要报名的排在最前，带倒计时。',
        '近期七天日程：开始时间与截止时间一目了然，做完的自动移到「已完成」。',
        '已读 / 未读、已完成标记在多设备间同步。',
        '网页、安卓 App、iPhone 主屏幕都能用；App 支持自动更新。',
    ]:
        bullet(doc, text)

    heading(doc, '五、给班委的说明', 1)
    for text in [
        '发布：班委登录后可直接发布；其他委员导入的通知进入「待审核」，由班委审核发布。',
        '导入：「群消息导入」支持粘贴 QQ / 微信群聊文本，也支持群聊截图识别文字；'
        '草稿可逐条修改字段，确认后加入待审核。',
        '存档与纠错：发错或过期的通知点「归档」，对同学不可见；需要时可在「已归档」里「撤销归档」重新发布。',
        '权限：班长、团支书可以管理所有通知；其他委员不能归档或修改班长、团支书发布的通知。',
        '日历订阅：同学们可以在「个人账号」里复制订阅链接，把班级日程加进手机日历。',
    ]:
        bullet(doc, text)

    heading(doc, '六、常见问题', 1)
    for text in [
        '收不到通知？依次检查：手机设置里允许通知、允许自启动、关闭电池优化，且不要从后台把 App 划掉。',
        '忘记密码？找班长在管理端重置，重置后请尽快改成自己的密码。',
        '换了手机？重新用学号登录，并在新手机上再打开一次通知开关即可（旧设备会失效）。',
        '消息发得晚或临时有变？以看板上的最新通知为准，旧通知过期后会自动置灰。',
    ]:
        bullet(doc, text)

    heading(doc, '七、版本与维护', 1)
    paragraph(doc, f'当前版本：网站 v{VERSION}（{UPDATED} 更新）· 安卓 App {VERSION}（versionCode {VERSION_CODE}）。'
                   '使用中遇到问题或想加功能，直接在群里说，或私聊班长。', size=10.5)
    paragraph(doc, '本说明不含任何账号密码；初始密码由班长单独分发，请注意不要转发给班级以外的人。', size=10)

    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, FILENAME)
    doc.save(path)
    return path


if __name__ == '__main__':
    target = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.expanduser('~'), 'OneDrive', 'Desktop')
    saved = build(target)
    print('已生成：' + saved)
